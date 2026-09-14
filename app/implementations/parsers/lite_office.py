"""DOCX and PPTX blocks via python-docx and python-pptx.

DOCX keeps Word's own heading styles, so the hierarchy is exact rather than
inferred. Body children are walked in document order so a table stays where the
author put it, between the paragraphs that introduce and follow it.

PPTX emits exactly one block group per slide, speaker notes included.
"""

from __future__ import annotations

import re

from app.interfaces.types import Block, BlockKind

_HEADING_STYLE = re.compile(r"^heading\s*(\d)$", re.IGNORECASE)
_LIST_STYLE = re.compile(r"(list|bullet)", re.IGNORECASE)


def _rows_to_markdown(rows: list[list[str]]) -> tuple[str, str | None]:
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return "", None
    width = max(len(r) for r in rows)
    padded = [r + [""] * (width - len(r)) for r in rows]
    header = "| " + " | ".join(padded[0]) + " |"
    separator = "| " + " | ".join(["---"] * width) + " |"
    body = ["| " + " | ".join(r) + " |" for r in padded[1:]]
    return "\n".join([header, separator, *body]), f"{header}\n{separator}"


def _docx_table(table) -> tuple[str, str | None]:
    rows = [
        [cell.text.replace("\n", " ").strip() for cell in row.cells]
        for row in table.rows
    ]
    return _rows_to_markdown(rows)


def parse_docx(file_path: str) -> tuple[list[Block], int | None, str]:
    from docx import Document
    from docx.table import Table

    document = Document(file_path)
    blocks: list[Block] = []
    markdown: list[str] = []
    section_path: list[str] = []
    levels: list[int] = []

    body = document.element.body
    tables = iter(document.tables)
    paragraphs = iter(document.paragraphs)

    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = next(paragraphs, None)
            if paragraph is None:
                continue
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "") if paragraph.style else ""
            heading = _HEADING_STYLE.match(style.strip()) or (
                style.strip().lower() == "title" and 1
            )
            if isinstance(heading, re.Match):
                level = int(heading.group(1))
            elif heading:
                level = 1
            else:
                level = None

            if level is not None:
                while levels and levels[-1] >= level:
                    levels.pop()
                    section_path.pop()
                levels.append(level)
                section_path.append(text)
                blocks.append(
                    Block(
                        kind=BlockKind.HEADING,
                        text=text,
                        section_path=list(section_path),
                        level=level,
                    )
                )
                markdown.append(f"{'#' * min(level, 6)} {text}")
                continue

            kind = BlockKind.LIST if _LIST_STYLE.search(style) else BlockKind.TEXT
            blocks.append(
                Block(kind=kind, text=text, section_path=list(section_path))
            )
            markdown.append(text)

        elif tag == "tbl":
            table = next(tables, None)
            if table is None:
                continue
            table_md, header = _docx_table(Table(child, document))
            if not table_md:
                continue
            blocks.append(
                Block(
                    kind=BlockKind.TABLE,
                    text=table_md,
                    section_path=list(section_path),
                    table_header=header,
                )
            )
            markdown.append(table_md)

    return blocks, None, "\n\n".join(markdown)


def _shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    lines = [p.text.strip() for p in shape.text_frame.paragraphs]
    return "\n".join(ln for ln in lines if ln)


def _pptx_table(shape) -> tuple[str, str | None]:
    rows = [
        [cell.text.replace("\n", " ").strip() for cell in row.cells]
        for row in shape.table.rows
    ]
    return _rows_to_markdown(rows)


def parse_pptx(file_path: str) -> tuple[list[Block], int, str]:
    from pptx import Presentation

    presentation = Presentation(file_path)
    blocks: list[Block] = []
    markdown: list[str] = []

    for index, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = (title_shape.text or "").strip() if title_shape is not None else ""
        section_path = [title] if title else [f"Slide {index}"]
        # `shapes.title` is a different wrapper object from the one iteration
        # yields, so identity does not hold -- match on the shape id instead.
        title_id = title_shape.shape_id if title_shape is not None else None

        if title:
            blocks.append(
                Block(
                    kind=BlockKind.HEADING,
                    text=title,
                    section_path=list(section_path),
                    level=1,
                    page=index,
                    slide_index=index,
                )
            )
            markdown.append(f"## {title}")

        for shape in slide.shapes:
            if title_id is not None and shape.shape_id == title_id:
                continue
            if getattr(shape, "has_table", False):
                table_md, header = _pptx_table(shape)
                if table_md:
                    blocks.append(
                        Block(
                            kind=BlockKind.TABLE,
                            text=table_md,
                            section_path=list(section_path),
                            page=index,
                            slide_index=index,
                            table_header=header,
                        )
                    )
                    markdown.append(table_md)
                continue
            text = _shape_text(shape)
            if text:
                blocks.append(
                    Block(
                        kind=BlockKind.TEXT,
                        text=text,
                        section_path=list(section_path),
                        page=index,
                        slide_index=index,
                    )
                )
                markdown.append(text)

        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                blocks.append(
                    Block(
                        kind=BlockKind.SLIDE_NOTES,
                        text=f"Speaker notes: {notes}",
                        section_path=list(section_path),
                        page=index,
                        slide_index=index,
                    )
                )
                markdown.append(f"> Speaker notes: {notes}")

    return blocks, len(presentation.slides), "\n\n".join(markdown)
