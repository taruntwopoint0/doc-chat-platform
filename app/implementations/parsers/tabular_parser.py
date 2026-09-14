"""Spreadsheets and delimited text.

Each row is serialised as `col: value | col: value` so a retrieved row carries
its own column names. That makes a row self-describing in isolation, which a
bare markdown table row is not -- and it removes the need to repeat a header
when a sheet spans many chunks.

Deliberately no pandas: this only ever walks rows and formats strings, and
pandas plus numpy costs over 100 MB of image for that. openpyxl reads XLSX in
streaming mode, the standard library reads CSV, and xlrd handles legacy XLS.
"""

from __future__ import annotations

import asyncio
import csv
from collections.abc import Iterator

from app.implementations.parsers import mime as mimes
from app.interfaces.parser import DocumentParser
from app.interfaces.types import Block, BlockKind, ParsedDocument

#: Guard against a stray column of 50k-character blobs blowing up a chunk.
_MAX_CELL_CHARS = 2000

Sheet = tuple[str, list[str], Iterator[list[str]]]


def _cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none"} else text


def _serialise_row(columns: list[str], row: list[str]) -> str:
    parts = []
    for column, value in zip(columns, row, strict=False):
        if not column or not value:
            continue
        if len(value) > _MAX_CELL_CHARS:
            value = value[:_MAX_CELL_CHARS] + "..."
        parts.append(f"{column}: {value}")
    return " | ".join(parts)


def _sniff_delimiter(path: str) -> str:
    with open(path, encoding="utf-8", errors="replace", newline="") as fh:
        sample = fh.read(8192)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _read_csv(path: str) -> list[Sheet]:
    delimiter = _sniff_delimiter(path)
    with open(path, encoding="utf-8-sig", errors="replace", newline="") as fh:
        reader = csv.reader(fh, delimiter=delimiter)
        rows = [[_cell(c) for c in row] for row in reader]
    if not rows:
        return []
    return [("", rows[0], iter(rows[1:]))]


def _read_xlsx(path: str) -> list[Sheet]:
    import openpyxl

    # read_only streams rows instead of building the whole sheet in memory,
    # which matters on a 2 GB instance with a large workbook.
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets: list[Sheet] = []
    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            continue
        columns = [_cell(c) for c in header]
        sheets.append(
            (worksheet.title, columns, ([_cell(c) for c in row] for row in rows))
        )
    return sheets


def _read_xls(path: str) -> list[Sheet]:
    import xlrd

    book = xlrd.open_workbook(path)
    sheets: list[Sheet] = []
    for worksheet in book.sheets():
        if worksheet.nrows == 0:
            continue
        columns = [_cell(worksheet.cell_value(0, c)) for c in range(worksheet.ncols)]
        rows = (
            [_cell(worksheet.cell_value(r, c)) for c in range(worksheet.ncols)]
            for r in range(1, worksheet.nrows)
        )
        sheets.append((worksheet.name, columns, rows))
    return sheets


def _load_sheets(path: str, mime_type: str) -> list[Sheet]:
    normalised = mimes.normalise(mime_type)
    if normalised == mimes.CSV:
        return _read_csv(path)
    if normalised == mimes.XLS:
        return _read_xls(path)
    return _read_xlsx(path)


def _sheets_to_blocks(sheets: list[Sheet]) -> tuple[list[Block], list[str], dict]:
    blocks: list[Block] = []
    markdown: list[str] = []
    counts: dict[str, int] = {}

    for index, (name, columns, rows) in enumerate(sheets, start=1):
        section_path = [name] if name else []
        if name:
            blocks.append(
                Block(
                    kind=BlockKind.HEADING,
                    text=name,
                    section_path=section_path,
                    level=1,
                )
            )
            markdown.append(f"## {name}")

        written = 0
        for row in rows:
            line = _serialise_row(columns, row)
            if not line:
                continue
            # One block per row: the chunker packs whole rows up to the token
            # target, so a row is never cut in half.
            blocks.append(
                Block(
                    kind=BlockKind.TABLE,
                    text=line,
                    section_path=list(section_path),
                    atomic=True,
                )
            )
            markdown.append(line)
            written += 1
        counts[name or f"sheet{index}"] = written

    return blocks, markdown, counts


class TabularParser(DocumentParser):
    # XLS is claimed by LegacyOfficeParser (spec routes legacy formats through
    # soffice); it delegates back here, and _load_sheets keeps the xlrd branch
    # as a native fallback for when soffice is not installed.
    mime_types = (mimes.CSV, mimes.XLSX)

    async def parse(self, file_path: str, mime_type: str) -> ParsedDocument:
        def work():
            sheets = _load_sheets(file_path, mime_type)
            return _sheets_to_blocks(sheets)

        blocks, markdown, counts = await asyncio.to_thread(work)
        return ParsedDocument(
            markdown="\n".join(markdown),
            blocks=blocks,
            page_count=len(counts) or None,
            metadata={
                "parser": "tabular",
                "mime_type": mime_type,
                "sheets": counts,
            },
        )
